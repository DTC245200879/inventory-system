"""Báo cáo: đọc số tiền bằng chữ (cho phiếu in) và xuất lịch sử giao dịch ra Excel."""
from io import BytesIO

DIGITS = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
GROUP_UNITS = ["", "nghìn", "triệu", "tỷ", "nghìn tỷ", "triệu tỷ"]


def _read_triple(n, full):
    """Đọc một nhóm 3 chữ số; full=True đọc cả 'không trăm', 'linh' (nhóm không đứng đầu)."""
    h, t, u = n // 100, (n // 10) % 10, n % 10
    words = []
    if full or h:
        words += [DIGITS[h], "trăm"]
    if t == 0:
        if u and (full or h):
            words.append("linh")
    elif t == 1:
        words.append("mười")
    else:
        words += [DIGITS[t], "mươi"]
    if u:
        if t >= 2 and u == 1:
            words.append("mốt")
        elif t >= 1 and u == 5:
            words.append("lăm")
        elif t >= 2 and u == 4:
            words.append("tư")
        else:
            words.append(DIGITS[u])
    return words


def so_thanh_chu(amount):
    """1250000 -> 'Một triệu hai trăm năm mươi nghìn đồng'."""
    n = int(round(float(amount or 0)))
    if n == 0:
        return "Không đồng"
    groups = []
    while n > 0:
        groups.append(n % 1000)
        n //= 1000
    words = []
    for i in range(len(groups) - 1, -1, -1):
        if groups[i] == 0:
            continue
        words += _read_triple(groups[i], full=(i != len(groups) - 1))
        if GROUP_UNITS[i]:
            words.append(GROUP_UNITS[i])
    text = " ".join(words)
    return text[0].upper() + text[1:] + " đồng"


def slip_code(tx):
    """Số phiếu theo kiểu chứng từ kho: PN-000012 (nhập), PX-000013 (xuất)."""
    return f"{'PN' if tx.type == 'IN' else 'PX'}-{tx.id:06d}"


def transactions_xlsx(items, period_text):
    """Tạo file .xlsx (bytes) từ danh sách giao dịch đã lọc."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Giao dịch"
    ws["A1"] = "Lịch sử nhập / xuất kho"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = period_text
    ws["A2"].font = Font(italic=True, color="64748B")
    headers = ["Số phiếu", "Thời gian", "Loại", "Mã SKU", "Sản phẩm", "ĐVT", "Số lượng",
               "Đơn giá (₫)", "Thành tiền (₫)", "Nhà cung cấp", "Ghi chú", "Người lập"]
    ws.append([])
    ws.append(headers)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(vertical="center")
    total_in = total_out = 0
    for t in items:
        amount = float(t.unit_price) * t.quantity
        ws.append([slip_code(t), t.created_at, "Nhập" if t.type == "IN" else "Xuất",
                   t.product.sku, t.product.name, t.product.unit,
                   t.quantity if t.type == "IN" else -t.quantity,
                   float(t.unit_price), amount,
                   t.supplier.name if t.supplier else "", t.note or "", t.created_by or ""])
        if t.type == "IN":
            total_in += amount
        else:
            total_out += amount
    for row in ws.iter_rows(min_row=head_row + 1, max_row=ws.max_row):
        row[1].number_format = "dd/mm/yyyy hh:mm"
        for c in (row[6], row[7], row[8]):
            c.number_format = "#,##0"
    ws.append([])
    ws.append(["", "", "", "", "Tổng giá trị nhập", "", "", "", total_in])
    ws.append(["", "", "", "", "Tổng giá trị xuất", "", "", "", total_out])
    for r in (ws.max_row - 1, ws.max_row):
        ws.cell(r, 5).font = Font(bold=True)
        ws.cell(r, 9).font = Font(bold=True)
        ws.cell(r, 9).number_format = "#,##0"
    widths = [12, 17, 8, 11, 34, 7, 10, 14, 16, 32, 30, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(headers))}{head_row + len(items)}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
